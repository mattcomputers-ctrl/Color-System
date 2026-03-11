"""Export utilities for PDF and Excel generation."""

import io
from datetime import datetime

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
)
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def generate_formula_pdf(formula_data: dict, target_data: dict = None) -> bytes:
    """Generate a PDF report for a formula.

    Args:
        formula_data: Dict with formula details and components.
        target_data: Optional dict with target color info.

    Returns:
        PDF file content as bytes.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter,
                            topMargin=0.75 * inch, bottomMargin=0.75 * inch)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'FormulaTitle', parent=styles['Title'], fontSize=16, spaceAfter=12
    )
    heading_style = ParagraphStyle(
        'FormulaHeading', parent=styles['Heading2'], fontSize=12, spaceAfter=6
    )

    elements = []

    # Title
    title = formula_data.get('title', 'Color Formula Report')
    elements.append(Paragraph(title, title_style))
    elements.append(Spacer(1, 12))

    # Target info
    if target_data:
        elements.append(Paragraph('Target Color', heading_style))
        target_info = [
            ['Identifier:', target_data.get('code', 'N/A')],
            ['Name:', target_data.get('name', 'N/A')],
        ]
        lab = target_data.get('lab_values')
        if lab:
            target_info.append(['Target L*a*b*:',
                                f"L*={lab['L']:.2f}  a*={lab['a']:.2f}  b*={lab['b']:.2f}"])

        t = Table(target_info, colWidths=[1.5 * inch, 4.5 * inch])
        t.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(t)
        elements.append(Spacer(1, 12))

    # Formula details
    elements.append(Paragraph('Formula Details', heading_style))
    details = [
        ['Series:', formula_data.get('series_name', 'N/A')],
        ['Version:', str(formula_data.get('version', 1))],
        ['Date:', formula_data.get('created_at', datetime.now().strftime('%Y-%m-%d %H:%M'))],
        ['Status:', formula_data.get('status', 'Generated')],
    ]
    pred_lab = formula_data.get('predicted_lab')
    if pred_lab:
        details.append(['Predicted L*a*b*:',
                         f"L*={pred_lab['L']:.2f}  a*={pred_lab['a']:.2f}  b*={pred_lab['b']:.2f}"])
    if formula_data.get('delta_e_2000') is not None:
        details.append(['ΔE*₀₀:', f"{formula_data['delta_e_2000']:.4f}"])
    if formula_data.get('delta_e_76') is not None:
        details.append(['ΔE*₇₆:', f"{formula_data['delta_e_76']:.4f}"])

    t = Table(details, colWidths=[1.5 * inch, 4.5 * inch])
    t.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(t)
    elements.append(Spacer(1, 16))

    # Components table
    elements.append(Paragraph('Formula Components', heading_style))
    components = formula_data.get('components', [])
    if components:
        header = ['Base Code', 'Base Name', 'Percentage (%)', 'Weight (g)']
        rows = [header]
        for comp in components:
            rows.append([
                comp.get('base_code', ''),
                comp.get('base_name', ''),
                f"{comp.get('percentage', 0):.2f}",
                f"{comp.get('weight_grams', 0):.2f}" if comp.get('weight_grams') else '—',
            ])
        # Total row
        total_pct = sum(c.get('percentage', 0) for c in components)
        total_wt = sum(c.get('weight_grams', 0) for c in components if c.get('weight_grams'))
        rows.append(['', 'TOTAL', f'{total_pct:.2f}', f'{total_wt:.2f}' if total_wt else '—'])

        t = Table(rows, colWidths=[1.2 * inch, 2.5 * inch, 1.2 * inch, 1.2 * inch])
        t.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2c3e50')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (2, 0), (-1, -1), 'RIGHT'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('LINEABOVE', (0, -1), (-1, -1), 1, colors.black),
        ]))
        elements.append(t)

    # Notes
    notes = formula_data.get('notes')
    if notes:
        elements.append(Spacer(1, 12))
        elements.append(Paragraph('Notes', heading_style))
        elements.append(Paragraph(notes, styles['Normal']))

    # Footer
    elements.append(Spacer(1, 24))
    elements.append(Paragraph(
        f'Generated by Color Formulation System — {datetime.now().strftime("%Y-%m-%d %H:%M")}',
        ParagraphStyle('Footer', parent=styles['Normal'], fontSize=8, textColor=colors.grey)
    ))

    doc.build(elements)
    return buffer.getvalue()


def generate_formula_excel(formulas: list[dict]) -> bytes:
    """Generate an Excel workbook with formula data.

    Args:
        formulas: List of formula dicts with components.

    Returns:
        Excel file content as bytes.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Formulas'

    # Styles
    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    # Headers
    headers = ['Target', 'Series', 'Version', 'Status',
               'Predicted L*', 'Predicted a*', 'Predicted b*',
               'ΔE*₀₀', 'ΔE*₇₆',
               'Base Code', 'Base Name', 'Percentage (%)', 'Weight (g)',
               'Date', 'Notes']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    # Data rows
    row = 2
    for formula in formulas:
        components = formula.get('components', [{}])
        for i, comp in enumerate(components):
            if i == 0:
                ws.cell(row=row, column=1, value=formula.get('target_code', ''))
                ws.cell(row=row, column=2, value=formula.get('series_name', ''))
                ws.cell(row=row, column=3, value=formula.get('version', ''))
                ws.cell(row=row, column=4, value=formula.get('status', ''))
                pred = formula.get('predicted_lab', {})
                ws.cell(row=row, column=5, value=pred.get('L'))
                ws.cell(row=row, column=6, value=pred.get('a'))
                ws.cell(row=row, column=7, value=pred.get('b'))
                ws.cell(row=row, column=8, value=formula.get('delta_e_2000'))
                ws.cell(row=row, column=9, value=formula.get('delta_e_76'))
                ws.cell(row=row, column=14, value=formula.get('created_at', ''))
                ws.cell(row=row, column=15, value=formula.get('notes', ''))

            ws.cell(row=row, column=10, value=comp.get('base_code', ''))
            ws.cell(row=row, column=11, value=comp.get('base_name', ''))
            ws.cell(row=row, column=12, value=comp.get('percentage'))
            ws.cell(row=row, column=13, value=comp.get('weight_grams'))

            for col in range(1, len(headers) + 1):
                ws.cell(row=row, column=col).border = thin_border
            row += 1

    # Auto-width columns
    for col in ws.columns:
        max_length = 0
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 30)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
